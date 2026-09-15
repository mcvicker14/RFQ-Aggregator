"""Plain-text extraction from uploaded solicitation documents, ahead of AI analysis."""
from pathlib import Path

import docx
from pypdf import PdfReader


class UnsupportedDocumentError(ValueError):
    pass


def extract_text(path: Path, content_type: str | None = None) -> str:
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages).strip()

    if suffix == ".docx":
        document = docx.Document(str(path))
        paragraphs = [p.text for p in document.paragraphs]
        tables_text = []
        for table in document.tables:
            for row in table.rows:
                tables_text.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(paragraphs + tables_text).strip()

    if suffix in (".txt", ".csv"):
        return path.read_text(errors="ignore").strip()

    raise UnsupportedDocumentError(
        f"AI analysis does not support '{suffix}' files yet. Supported: PDF, DOCX, TXT, CSV."
    )
